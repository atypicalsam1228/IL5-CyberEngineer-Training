"""Shared document builder for IL5 cyber-engineering training modules.

Writes a .docx and a parallel .md from the same calls, and counts callouts so
each generator can print the cross-reference summary the format requires.
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")

CODE_SHADE = "F3F3F3"
OUTPUT_SHADE = "F0F8F0"


def _shade(paragraph, color: str) -> None:
    props = paragraph._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "solid")
    shd.set(qn("w:color"), color)
    shd.set(qn("w:fill"), color)
    props.append(shd)


class DocBuilder:
    def __init__(self) -> None:
        self.doc = Document()
        self.md_lines: list[str] = []
        self.counts: Counter[str] = Counter()
        self._style_headings()

    def _style_headings(self) -> None:
        self.doc.styles["Title"].font.size = Pt(24)
        self.doc.styles["Title"].font.color.rgb = RGBColor(0, 51, 102)
        for level, size, color in (
            ("Heading 1", 18, RGBColor(0, 51, 102)),
            ("Heading 2", 14, RGBColor(0, 80, 140)),
            ("Heading 3", 12, RGBColor(0, 100, 160)),
        ):
            style = self.doc.styles[level]
            style.font.size = Pt(size)
            style.font.color.rgb = color

    def _md(self, text: str = "") -> None:
        self.md_lines.append(text)

    def title(self, text: str) -> None:
        self.doc.add_heading(text, level=0)
        self._md(f"# {text}\n")

    def h1(self, text: str) -> None:
        self.doc.add_heading(text, level=1)
        self._md(f"## {text}\n")

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)
        self._md(f"### {text}\n")

    def h3(self, text: str) -> None:
        self.doc.add_heading(text, level=3)
        self._md(f"#### {text}\n")

    def para(self, text: str) -> None:
        paragraph = self.doc.add_paragraph(text)
        for run in paragraph.runs:
            run.font.size = Pt(11)
        self._md(f"{text}\n")

    def bullet(self, text: str) -> None:
        paragraph = self.doc.add_paragraph(text, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.size = Pt(11)
        self._md(f"- {text}")

    def numbered(self, text: str) -> None:
        paragraph = self.doc.add_paragraph(text, style="List Number")
        for run in paragraph.runs:
            run.font.size = Pt(11)
        self._md(f"1. {text}")

    def _callout(self, key: str, emoji: str, text: str, italic: bool, color: RGBColor | None, size: int = 11) -> None:
        self.counts[key] += 1
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run(f"{emoji} {text}")
        run.font.size = Pt(size)
        run.font.italic = italic
        if color is not None:
            run.font.color.rgb = color
        self._md(f"> {emoji} {text}\n")

    def note(self, text: str) -> None:
        self._callout("note", "\U0001f4cc", text, True, None)

    def security(self, text: str) -> None:
        self._callout("security", "\U0001f512", text, True, RGBColor(150, 0, 0))

    def tip(self, text: str) -> None:
        self._callout("tip", "\U0001f4a1", text, True, RGBColor(0, 100, 60))

    def production(self, text: str) -> None:
        self._callout("production", "\U0001f3ed", f"IN PRODUCTION: {text}", True, RGBColor(100, 100, 100), 10)

    def incident(self, text: str) -> None:
        self._callout("incident", "\U0001f4f0", f"REAL-WORLD: {text}", True, RGBColor(140, 80, 0), 10)

    def action(self, text: str) -> None:
        self._callout("action", "▶", f"YOUR TURN: {text}", False, RGBColor(0, 80, 160))

    def checkpoint(self, text: str) -> None:
        self._callout("checkpoint", "✅", f"CHECKPOINT: {text}", False, RGBColor(0, 120, 60))

    def code(self, lines: str, lang: str = "bash") -> None:
        for line in lines.split("\n"):
            paragraph = self.doc.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            paragraph.paragraph_format.space_after = Pt(0)
            _shade(paragraph, CODE_SHADE)
        self.doc.add_paragraph("")
        self._md(f"```{lang}\n{lines}\n```\n")

    def output(self, lines: str) -> None:
        label = self.doc.add_paragraph()
        run = label.add_run("Expected output:")
        run.bold = True
        run.font.size = Pt(10)
        for line in lines.split("\n"):
            paragraph = self.doc.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0, 100, 0)
            paragraph.paragraph_format.space_after = Pt(0)
            _shade(paragraph, OUTPUT_SHADE)
        self.doc.add_paragraph("")
        self._md(f"**Expected output:**\n\n```\n{lines}\n```\n")

    def image(self, path: Path, alt: str, ascii_text: str = "", width_in: float = 6.3) -> None:
        """Insert a PNG into the .docx; the .md gets the ASCII diagram (if given) plus an image link."""
        self.doc.add_picture(str(path), width=Inches(width_in))
        self.doc.add_paragraph("")
        if ascii_text:
            self._md("```text\n" + ascii_text + "\n```\n")
        self._md(f"![{alt}](img/{Path(path).name})\n")

    def table(self, headers: list[str], rows: list[list[str]]) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(10)
        for row_data in rows:
            row = table.add_row()
            for index, value in enumerate(row_data):
                row.cells[index].text = value
                for run in row.cells[index].paragraphs[0].runs:
                    run.font.size = Pt(10)
        self.doc.add_paragraph("")
        self._md("| " + " | ".join(headers) + " |")
        self._md("| " + " | ".join("---" for _ in headers) + " |")
        for row_data in rows:
            self._md("| " + " | ".join(cell.replace("\n", " ") for cell in row_data) + " |")
        self._md()

    def _link_cell(self, cell, text: str, url: str) -> None:
        part = self.doc.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "20")
        for el in (color, underline, size):
            props.append(el)
        run.append(props)
        label = OxmlElement("w:t")
        label.text = text
        run.append(label)
        link.append(run)
        cell.paragraphs[0]._p.append(link)

    def tool_table(self, rows: list[tuple[str, str, str, str]]) -> None:
        """Table of (name, url, what it is, used here); the name is a live link in the .docx."""
        headers = ["Tool (overview link)", "What it is", "Used in this lab"]
        table = self.doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(10)
        for name, url, what, use in rows:
            row = table.add_row()
            self._link_cell(row.cells[0], name, url)
            for index, value in ((1, what), (2, use)):
                row.cells[index].text = value
                for run in row.cells[index].paragraphs[0].runs:
                    run.font.size = Pt(9)
        self.doc.add_paragraph("")
        self._md("| " + " | ".join(headers) + " |")
        self._md("| --- | --- | --- |")
        for name, url, what, use in rows:
            self._md(f"| [{name}]({url}) | {what} | {use} |")
        self._md()

    def save(self, directory: Path, stem: str) -> tuple[Path, Path]:
        directory.mkdir(parents=True, exist_ok=True)
        docx_path = directory / f"{stem}.docx"
        md_path = directory / f"{stem}.md"
        self.doc.save(docx_path)
        md_path.write_text("\n".join(self.md_lines), encoding="utf-8")
        return docx_path, md_path
